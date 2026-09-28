import hashlib
import json
import re
import unicodedata
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile, status
from sqlalchemy import desc, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session, selectinload

from app.db.session import get_db
from app.models.user import User
from app.models.ws3_order import WS3ImportBatch, WS3ImportRow, WS3Order, WS3OrderRoll
from app.schemas.ws3_order import WS3ImportBatchCreate, WS3ImportBatchDetailResponse, WS3ImportBatchResponse, WS3ImportRowResponse, WS3ImportWarehouseResponse, WS3OrderImport, WS3OrderResponse, WS3OrderRollResponse, WS3OrderStatusUpdate, WS3WarehouseRowUpdate
from app.security.auth import get_current_user, role_codes

router = APIRouter(prefix="/ws3/orders", tags=["ws3-orders"])

FIELD_ALIASES = {
    "roll_id": ("roll id", "roll", "cuộn", "mã cuộn", "qr"),
    "item_code": ("item code", "item", "mã hàng"),
    "item_name": ("item name", "fabric", "fabric name", "tên vải"),
    "lot_no": ("lot no", "lot", "số lot"),
    "machine_no": ("m/c no", "machine", "loom", "loom no", "số máy"),
    "length_meters": ("length", "meters", "met", "mét", "sản lượng"),
    "shift": ("shift", "ca"),
    "worker": ("worker", "operator", "người vận hành"),
    "remarks": ("remarks", "remark", "note", "ghi chú", "sop"),
}

MES_HEADERS = {
    "DATE": ("date", "ngay", "ngày", "date production"),
    "SHIFT": ("shift", "ca"),
    "CUSTOMER": ("customer", "khach hang", "khách hàng"),
    "PO #": ("po", "po #", "po no"),
    "ORDER #": ("order", "order #", "order no"),
    "ITEM NO": ("item no", "item code", "item"),
    "FABRIC": ("fabric", "item name", "ten vai", "tên vải"),
    "STYLE #": ("style", "style #"),
    "COLOR": ("color", "mau", "màu"),
    "SOP #": ("sop", "sop #"),
    "LOT #": ("lot", "lot #", "lot no"),
    "MACHINE": ("machine", "m/c", "loom", "loom no"),
    "QTY UNIT": ("qty unit", "unit"),
    "ROLL": ("roll", "roll no", "cuon", "cuộn"),
    "KGS": ("kgs", "kg"),
    "YDS": ("yds", "yard"),
    "MTS": ("mts", "met", "meter", "meters", "length"),
    "WORKER": ("worker", "operator", "nguoi van hanh", "người vận hành"),
}

SOURCE_KEY_ALIASES = {
    "date": ("DATE", "DATE PRODUCTION"),
    "shift": ("SHIFT",),
    "machine": ("MACHINE", "M/C NO", "LOOM", "LOOM NO"),
    "sop": ("SOP #", "SOP"),
    "lot": ("LOT #", "LOT", "LOT NO"),
    "mts": ("MTS", "MET", "METER", "METERS", "LENGTH"),
}


def _clean_header(value: object) -> str:
    text = unicodedata.normalize("NFKD", str(value or "")).encode("ascii", "ignore").decode().strip().lower()
    text = re.sub(r"\s+", " ", text)
    return text


def _canonical_header(value: object) -> str:
    cleaned = _clean_header(value)
    for canonical, aliases in MES_HEADERS.items():
        if cleaned in {_clean_header(alias) for alias in aliases}:
            return canonical
    return str(value).strip() if str(value or "").strip() else ""


def _date_like(value: object) -> bool:
    text = str(value or "").strip()
    return isinstance(value, (datetime,)) or bool(re.search(r"\d{1,4}[-/]\d{1,2}[-/]\d{1,4}", text)) or (text.isdigit() and 20000 <= int(text) <= 60000)


def _normalise_mts(value: object) -> str:
    text = str(value or "").strip().replace(",", "")
    try:
        number = Decimal(text)
        return format(number.normalize(), "f")
    except (InvalidOperation, ValueError):
        return re.sub(r"\s+", " ", text.upper())


def _source_key(row: dict[str, str]) -> str | None:
    values: list[str] = []
    normalized = {_clean_header(column): value for column, value in row.items()}
    for field, aliases in SOURCE_KEY_ALIASES.items():
        value = next((normalized.get(_clean_header(alias), "") for alias in aliases if normalized.get(_clean_header(alias), "")), "")
        if not value:
            return None
        value = _normalise_mts(value) if field == "mts" else value.strip().upper()
        if field == "date":
            value = value[:10].replace("/", "-")
        values.append(re.sub(r"\s+", " ", value))
    return "|".join(values)


def _fallback_source_key(row: dict[str, str]) -> str:
    """Fingerprint every source field when the MES business key is incomplete."""
    values = [
        (_clean_header(column), re.sub(r"\s+", " ", str(value or "").strip()).upper())
        for column, value in row.items()
    ]
    payload = json.dumps(sorted(values), ensure_ascii=False, separators=(",", ":"))
    return f"FALLBACK:{hashlib.sha256(payload.encode('utf-8')).hexdigest()}"


def _dedupe_key(row: dict[str, str]) -> str:
    return _source_key(row) or _fallback_source_key(row)


def _existing_source_keys(db: Session, model: type[WS3OrderRoll] | type[WS3ImportRow]) -> set[str]:
    keys = set(db.scalars(select(model.source_key).where(model.source_key.is_not(None))).all())
    legacy_rows = db.scalars(select(model.raw_data_json).where(model.source_key.is_(None))).all()
    for raw_data in legacy_rows:
        try:
            keys.add(_dedupe_key(json.loads(raw_data or "{}")))
        except (TypeError, json.JSONDecodeError):
            continue
    return keys


def _dedupe_rows(rows: list[dict[str, str]], existing_keys: set[str]) -> tuple[list[int], set[int], set[int]]:
    seen = set(existing_keys)
    new_indexes: list[int] = []
    duplicate_indexes: set[int] = set()
    incomplete_indexes: set[int] = set()
    for index, row in enumerate(rows):
        primary_key = _source_key(row)
        key = primary_key or _fallback_source_key(row)
        if primary_key is None:
            incomplete_indexes.add(index)
        if key in seen:
            duplicate_indexes.add(index)
        else:
            seen.add(key)
            if primary_key is not None:
                new_indexes.append(index)
    return new_indexes, duplicate_indexes, incomplete_indexes


def _normalise_mes_values(values: list[list[object]]) -> tuple[list[str], list[dict[str, str]], int]:
    """Find the MES header row and return clean rows while preserving source fields."""
    header_index = -1
    best_score = 0
    for index, row in enumerate(values[:30]):
        score = sum(1 for value in row if _canonical_header(value) in MES_HEADERS)
        if score > best_score:
            best_score = score
            header_index = index
    if header_index < 0 or best_score < 3:
        raise HTTPException(status_code=422, detail="Could not find the MES header row")

    raw_headers = values[header_index]
    possible_subheaders = values[header_index + 1] if header_index + 1 < len(values) else []
    subheader_score = sum(1 for value in possible_subheaders if _canonical_header(value) in MES_HEADERS)
    has_subheaders = subheader_score >= 3
    data_start = header_index + 2 if has_subheaders else header_index + 1
    columns: list[str] = []
    used: set[str] = set()
    source_indexes: list[int] = []
    for index, value in enumerate(raw_headers):
        parent = _clean_header(value)
        child_value = possible_subheaders[index] if has_subheaders and index < len(possible_subheaders) else ""
        child = _canonical_header(child_value)
        # MES exports merge "PRODUCTION QTY" over four child columns. The
        # child headers are the real fields and must never be collapsed into
        # one production-quantity column.
        canonical = child if child in {"ROLL", "KGS", "YDS", "MTS"} else _canonical_header(value)
        if not canonical and child_value:
            canonical = str(child_value).strip()
        if not canonical:
            continue
        column = canonical
        suffix = 2
        while column in used:
            column = f"{canonical} ({suffix})"
            suffix += 1
        columns.append(column)
        used.add(column)
        source_indexes.append(index)

    rows: list[dict[str, str]] = []
    ignored = 0
    for raw_row in values[data_start:]:
        row = {column: "" if raw_row[index] is None else str(raw_row[index]).strip() for column, index in zip(columns, source_indexes)}
        if not any(row.values()):
            continue
        if sum(1 for value in raw_row if _canonical_header(value) in MES_HEADERS) >= 3:
            ignored += 1
            continue
        if "DATE" in row and row["DATE"] and not _date_like(row["DATE"]):
            ignored += 1
            continue
        rows.append(row)
    return columns, rows, ignored


def _require_supervisor(user: User) -> None:
    if not role_codes(user).intersection({"ADMIN", "SUPERVISOR", "PRODUCTION_MANAGER"}):
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Supervisor access is required")


def _mapped(row: dict[str, str], mapping: dict[str, str], field: str) -> str | None:
    source_column = mapping.get(field)
    if not source_column:
        return None
    value = row.get(source_column)
    return value.strip() if isinstance(value, str) and value.strip() else None


def _decimal(value: str | None) -> Decimal | None:
    if not value:
        return None
    normalized = value.replace(",", "").strip()
    try:
        return Decimal(normalized)
    except InvalidOperation:
        return None


def _response(order: WS3Order, *, duplicate_count: int = 0, new_count: int | None = None, incomplete_key_count: int = 0) -> WS3OrderResponse:
    return WS3OrderResponse(
        id=order.id,
        order_no=order.order_no,
        source_filename=order.source_filename,
        source_format=order.source_format,
        columns=order.source_columns,
        mapping=order.selected_mapping,
        status=order.status,
        created_at=order.created_at,
        confirmed_at=order.confirmed_at,
        rolls=[WS3OrderRollResponse(
            id=roll.id,
            source_row_index=roll.source_row_index,
            source_key=roll.source_key,
            raw_data=json.loads(roll.raw_data_json or "{}"),
            roll_id=roll.roll_id,
            item_code=roll.item_code,
            item_name=roll.item_name,
            lot_no=roll.lot_no,
            machine_no=roll.machine_no,
            length_meters=float(roll.length_meters) if roll.length_meters is not None else None,
            shift=roll.shift,
            worker=roll.worker,
            remarks=roll.remarks,
            status=roll.status,
        ) for roll in order.rolls],
        duplicate_count=duplicate_count,
        new_count=len(order.rolls) if new_count is None else new_count,
        incomplete_key_count=incomplete_key_count,
    )


def _import_response(batch: WS3ImportBatch, *, duplicate_count: int | None = None, incomplete_key_count: int | None = None) -> WS3ImportBatchResponse:
    available_count = sum(1 for row in batch.rows if row.status == "AVAILABLE")
    duplicate_count = batch.duplicate_count if duplicate_count is None and batch.duplicate_count is not None else duplicate_count
    incomplete_key_count = batch.incomplete_key_count if incomplete_key_count is None and batch.incomplete_key_count is not None else incomplete_key_count
    duplicate_count = sum(1 for row in batch.rows if row.status == "DUPLICATE") if duplicate_count is None else duplicate_count
    incomplete_key_count = sum(1 for row in batch.rows if row.status == "INCOMPLETE") if incomplete_key_count is None else incomplete_key_count
    return WS3ImportBatchResponse(
        id=batch.id,
        source_filename=batch.source_filename,
        source_format=batch.source_format,
        created_at=batch.created_at,
        row_count=batch.source_row_count if batch.source_row_count is not None else len(batch.rows),
        available_count=available_count,
        duplicate_count=duplicate_count,
        incomplete_key_count=incomplete_key_count,
    )


def _import_detail_response(batch: WS3ImportBatch) -> WS3ImportBatchDetailResponse:
    summary = _import_response(batch)
    return WS3ImportBatchDetailResponse(
        **summary.model_dump(),
        columns=batch.source_columns,
        mapping=batch.selected_mapping,
        rows=[WS3ImportRowResponse(
            id=row.id,
            source_row_index=row.source_row_index,
            source_key=row.source_key,
            raw_data=json.loads(row.raw_data_json or "{}"),
            status=row.status,
            order_id=row.order_id,
        ) for row in batch.rows],
    )


def _unique_columns(values: list[object]) -> list[str]:
    columns: list[str] = []
    used: set[str] = set()
    for index, value in enumerate(values):
        base = str(value).strip() if value is not None else ""
        base = base or f"Column {index + 1}"
        column = base
        suffix = 2
        while column in used:
            column = f"{base} ({suffix})"
            suffix += 1
        columns.append(column)
        used.add(column)
    return columns


@router.post("/parse-file")
async def parse_source_file(file: UploadFile = File(...), db: Session = Depends(get_db), user: User = Depends(get_current_user)) -> dict[str, object]:
    _require_supervisor(user)
    filename = file.filename or "source"
    suffix = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    content = await file.read()
    if suffix in {"csv", "tsv", "txt"}:
        return {"filename": filename, "format": "FILE", "text": content.decode("utf-8-sig", errors="replace")}
    if suffix == "xls":
        import xlrd
        from xlrd.xldate import xldate_as_datetime

        workbook = xlrd.open_workbook(file_contents=content, on_demand=True)
        sheet = workbook.sheet_by_index(0)
        values = []
        for row in range(sheet.nrows):
            current: list[object] = []
            for column in range(sheet.ncols):
                cell = sheet.cell(row, column)
                current.append(xldate_as_datetime(cell.value, workbook.datemode) if cell.ctype == xlrd.XL_CELL_DATE else cell.value)
            values.append(current)
    elif suffix == "xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(filename=__import__("io").BytesIO(content), read_only=True, data_only=True)
        sheet = workbook.active
        values = [list(row) for row in sheet.iter_rows(values_only=True)]
    else:
        raise HTTPException(status_code=415, detail="Use an XLS, XLSX, CSV or TSV file")
    values = [row for row in values if any(value not in (None, "") for value in row)]
    if len(values) < 2:
        raise HTTPException(status_code=422, detail="The source file must contain a header and at least one data row")
    columns, rows, ignored_rows = _normalise_mes_values(values)
    existing_keys = _existing_source_keys(db, WS3OrderRoll)
    existing_keys.update(_existing_source_keys(db, WS3ImportRow))
    new_indexes, duplicate_indexes, incomplete_indexes = _dedupe_rows(rows, existing_keys)
    return {
        "filename": filename,
        "format": "FILE",
        "columns": columns,
        "rows": rows,
        "ignored_rows": ignored_rows,
        "normalization_message": f"Đã chuẩn hóa {len(rows):,} dòng dữ liệu MES; bỏ qua {ignored_rows:,} dòng tiêu đề/mẫu không hợp lệ.",
        "normalization_message_en": f"Normalized {len(rows):,} MES rows; skipped {ignored_rows:,} header/sample rows.",
        "new_row_indexes": new_indexes,
        "duplicate_row_indexes": sorted(duplicate_indexes),
        "incomplete_key_indexes": sorted(incomplete_indexes),
        "duplicate_count": len(duplicate_indexes),
        "new_count": len(new_indexes),
        "incomplete_key_count": len(incomplete_indexes),
    }


@router.post("/imports", response_model=WS3ImportBatchResponse, status_code=status.HTTP_201_CREATED)
def create_import_batch(
    payload: WS3ImportBatchCreate,
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
) -> WS3ImportBatchResponse:
    _require_supervisor(user)
    existing_keys = _existing_source_keys(db, WS3OrderRoll)
    existing_keys.update(_existing_source_keys(db, WS3ImportRow))
    selected = set(payload.selected_row_indexes if payload.selected_row_indexes is not None else range(len(payload.rows)))
    batch = WS3ImportBatch(
        source_filename=payload.source_filename,
        source_format=payload.source_format,
        source_columns_json=json.dumps(payload.columns, ensure_ascii=False),
        selected_mapping_json=json.dumps(payload.mapping, ensure_ascii=False),
        source_row_count=len(selected),
        created_by_id=user.id,
    )
    seen = set(existing_keys)
    duplicate_count = 0
    incomplete_key_count = 0
    for source_index, row in enumerate(payload.rows):
        if source_index not in selected:
            continue
        primary_key = _source_key(row)
        source_key = primary_key or _fallback_source_key(row)
        if primary_key is None:
            incomplete_key_count += 1
        if source_key in seen:
            duplicate_count += 1
            continue
        seen.add(source_key)
        row_status = "INCOMPLETE" if primary_key is None else "AVAILABLE"
        batch.rows.append(WS3ImportRow(
            source_row_index=source_index,
            source_key=source_key,
            raw_data_json=json.dumps(row, ensure_ascii=False),
            status=row_status,
        ))
    batch.duplicate_count = duplicate_count
    batch.incomplete_key_count = incomplete_key_count
    db.add(batch)
    db.commit()
    db.refresh(batch)
    return _import_response(batch, duplicate_count=duplicate_count, incomplete_key_count=incomplete_key_count)


@router.get("/imports", response_model=list[WS3ImportBatchResponse])
def list_import_batches(db: Session = Depends(get_db), user: User = Depends(get_current_user)) -> list[WS3ImportBatchResponse]:
    _require_supervisor(user)
    statement = select(WS3ImportBatch).options(selectinload(WS3ImportBatch.rows)).order_by(desc(WS3ImportBatch.created_at)).limit(100)
    return [_import_response(batch) for batch in db.scalars(statement).unique().all()]


@router.get("/imports/data", response_model=WS3ImportWarehouseResponse)
def get_import_warehouse(db: Session = Depends(get_db), user: User = Depends(get_current_user)) -> WS3ImportWarehouseResponse:
    _require_supervisor(user)
    batches = db.scalars(select(WS3ImportBatch).options(selectinload(WS3ImportBatch.rows)).order_by(WS3ImportBatch.created_at)).unique().all()
    columns: list[str] = []
    seen_columns: set[str] = set()
    rows: list[WS3ImportRowResponse] = []
    for batch in batches:
        for column in batch.source_columns:
            if column not in seen_columns:
                columns.append(column)
                seen_columns.add(column)
        for row in batch.rows:
            if row.status in {"DUPLICATE", "ARCHIVED"}:
                continue
            rows.append(WS3ImportRowResponse(
                id=row.id,
                source_row_index=row.source_row_index,
                source_key=row.source_key,
                raw_data=json.loads(row.raw_data_json or "{}"),
                status=row.status,
                order_id=row.order_id,
            ))
    return WS3ImportWarehouseResponse(columns=columns, rows=rows)


@router.patch("/imports/data/{row_id}", response_model=WS3ImportRowResponse)
def update_warehouse_row(row_id: int, payload: WS3WarehouseRowUpdate, db: Session = Depends(get_db), user: User = Depends(get_current_user)) -> WS3ImportRowResponse:
    _require_supervisor(user)
    row = db.scalar(select(WS3ImportRow).where(WS3ImportRow.id == row_id))
    if row is None or row.status == "ARCHIVED":
        raise HTTPException(status_code=404, detail="Warehouse row was not found")
    if row.status == "ASSIGNED":
        raise HTTPException(status_code=409, detail="A row already assigned to an order cannot be edited")
    source_key = _dedupe_key(payload.raw_data)
    other_import_keys = set(db.scalars(select(WS3ImportRow.source_key).where(WS3ImportRow.id != row_id, WS3ImportRow.source_key.is_not(None))).all())
    order_keys = _existing_source_keys(db, WS3OrderRoll)
    if source_key in other_import_keys or source_key in order_keys:
        raise HTTPException(status_code=409, detail="This change matches data already in the warehouse")
    row.raw_data_json = json.dumps(payload.raw_data, ensure_ascii=False)
    row.source_key = source_key
    row.status = "INCOMPLETE" if _source_key(payload.raw_data) is None else "AVAILABLE"
    db.commit()
    db.refresh(row)
    return WS3ImportRowResponse(id=row.id, source_row_index=row.source_row_index, source_key=row.source_key, raw_data=json.loads(row.raw_data_json), status=row.status, order_id=row.order_id)


@router.delete("/imports/data/{row_id}", status_code=status.HTTP_204_NO_CONTENT)
def archive_warehouse_row(row_id: int, db: Session = Depends(get_db), user: User = Depends(get_current_user)) -> None:
    _require_supervisor(user)
    row = db.scalar(select(WS3ImportRow).where(WS3ImportRow.id == row_id))
    if row is None or row.status == "ARCHIVED":
        raise HTTPException(status_code=404, detail="Warehouse row was not found")
    if row.status == "ASSIGNED":
        raise HTTPException(status_code=409, detail="A row already assigned to an order cannot be deleted")
    row.status = "ARCHIVED"
    db.commit()


@router.get("/imports/{batch_id}", response_model=WS3ImportBatchDetailResponse)
def get_import_batch(batch_id: int, db: Session = Depends(get_db), user: User = Depends(get_current_user)) -> WS3ImportBatchDetailResponse:
    _require_supervisor(user)
    batch = db.scalar(select(WS3ImportBatch).options(selectinload(WS3ImportBatch.rows)).where(WS3ImportBatch.id == batch_id))
    if batch is None:
        raise HTTPException(status_code=404, detail="MES import batch was not found")
    return _import_detail_response(batch)


@router.post("", response_model=WS3OrderResponse, status_code=status.HTTP_201_CREATED)
def create_order(
    payload: WS3OrderImport,
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
) -> WS3OrderResponse:
    _require_supervisor(user)
    import_rows_for_order: list[WS3ImportRow] = []
    if payload.import_row_ids is not None:
        selected_ids = set(payload.import_row_ids)
        import_rows_for_order = db.scalars(
            select(WS3ImportRow).where(WS3ImportRow.id.in_(selected_ids), WS3ImportRow.status == "AVAILABLE")
        ).all()
        if len(import_rows_for_order) != len(selected_ids):
            raise HTTPException(status_code=409, detail="One or more selected warehouse rows are no longer available")
        source_rows = [(row.id, json.loads(row.raw_data_json or "{}")) for row in import_rows_for_order]
    elif payload.import_batch_id is not None:
        batch = db.scalar(select(WS3ImportBatch).options(selectinload(WS3ImportBatch.rows)).where(WS3ImportBatch.id == payload.import_batch_id))
        if batch is None:
            raise HTTPException(status_code=404, detail="MES import batch was not found")
        batch_rows = {row.source_row_index: row for row in batch.rows if row.status == "AVAILABLE"}
        selected_indexes = set(payload.selected_row_indexes if payload.selected_row_indexes is not None else batch_rows.keys())
        source_rows = [(index, json.loads(row.raw_data_json or "{}")) for index, row in batch_rows.items() if index in selected_indexes]
        import_rows_for_order = [row for index, row in batch_rows.items() if index in selected_indexes]
        payload.mapping = payload.mapping or batch.selected_mapping
    else:
        source_rows = [(index, row) for index, row in enumerate(payload.rows) if index in set(payload.selected_row_indexes if payload.selected_row_indexes is not None else range(len(payload.rows)))]
    rows = source_rows
    if not rows:
        raise HTTPException(status_code=422, detail="Select at least one source row")

    existing_keys = set(db.scalars(select(WS3OrderRoll.source_key).where(WS3OrderRoll.source_key.is_not(None))).all())
    import_key_query = select(WS3ImportRow.source_key).where(WS3ImportRow.source_key.is_not(None))
    if payload.import_row_ids is not None:
        import_key_query = import_key_query.where(WS3ImportRow.id.not_in(payload.import_row_ids))
    elif payload.import_batch_id is not None:
        import_key_query = import_key_query.where(WS3ImportRow.batch_id != payload.import_batch_id)
    existing_keys.update(db.scalars(import_key_query).all())
    seen_keys = set(existing_keys)
    unique_rows: list[tuple[int, dict[str, str], str]] = []
    duplicate_count = 0
    incomplete_key_count = 0
    for source_index, row in rows:
        key = _source_key(row)
        if key is None:
            incomplete_key_count += 1
        elif key in seen_keys:
            duplicate_count += 1
        else:
            seen_keys.add(key)
            unique_rows.append((source_index, row, key))
    if not unique_rows:
        raise HTTPException(status_code=409, detail=f"No new rows to import. Duplicates: {duplicate_count}; incomplete keys: {incomplete_key_count}.")

    order = WS3Order(
        order_no=f"WS3-{datetime.now(timezone.utc):%Y%m%d%H%M%S}",
        source_filename=payload.source_filename,
        source_format=payload.source_format,
        source_columns_json=json.dumps(payload.columns, ensure_ascii=False),
        selected_mapping_json=json.dumps(payload.mapping, ensure_ascii=False),
        source_rows_json=json.dumps(payload.rows, ensure_ascii=False),
        created_by_id=user.id,
    )
    for source_index, row, source_key in unique_rows:
        order.rolls.append(WS3OrderRoll(
            source_row_index=source_index,
            source_key=source_key,
            raw_data_json=json.dumps(row, ensure_ascii=False),
            roll_id=_mapped(row, payload.mapping, "roll_id"),
            item_code=_mapped(row, payload.mapping, "item_code"),
            item_name=_mapped(row, payload.mapping, "item_name"),
            lot_no=_mapped(row, payload.mapping, "lot_no"),
            machine_no=_mapped(row, payload.mapping, "machine_no"),
            length_meters=_decimal(_mapped(row, payload.mapping, "length_meters")),
            shift=_mapped(row, payload.mapping, "shift"),
            worker=_mapped(row, payload.mapping, "worker"),
            remarks=_mapped(row, payload.mapping, "remarks"),
        ))
    db.add(order)
    db.flush()
    for import_row in import_rows_for_order:
        import_row.status = "ASSIGNED"
        import_row.order_id = order.id
    try:
        db.commit()
    except IntegrityError as error:
        db.rollback()
        raise HTTPException(status_code=409, detail="One or more MES source keys were imported by another user. Please upload the file again to refresh duplicate counts.") from error
    db.refresh(order)
    return _response(order, duplicate_count=duplicate_count, new_count=len(unique_rows), incomplete_key_count=incomplete_key_count)


@router.get("", response_model=list[WS3OrderResponse])
def list_orders(db: Session = Depends(get_db), user: User = Depends(get_current_user)) -> list[WS3OrderResponse]:
    statement = select(WS3Order).options(selectinload(WS3Order.rolls)).order_by(desc(WS3Order.created_at)).limit(100)
    return [_response(order) for order in db.scalars(statement).unique().all()]


@router.get("/{order_id}", response_model=WS3OrderResponse)
def get_order(order_id: int, db: Session = Depends(get_db), user: User = Depends(get_current_user)) -> WS3OrderResponse:
    order = db.scalar(select(WS3Order).options(selectinload(WS3Order.rolls)).where(WS3Order.id == order_id))
    if order is None:
        raise HTTPException(status_code=404, detail="WS3 order was not found")
    return _response(order)


@router.patch("/{order_id}/status", response_model=WS3OrderResponse)
def update_order_status(order_id: int, payload: WS3OrderStatusUpdate, db: Session = Depends(get_db), user: User = Depends(get_current_user)) -> WS3OrderResponse:
    _require_supervisor(user)
    order = db.scalar(select(WS3Order).options(selectinload(WS3Order.rolls)).where(WS3Order.id == order_id))
    if order is None:
        raise HTTPException(status_code=404, detail="WS3 order was not found")
    order.status = payload.status
    order.confirmed_at = datetime.now(timezone.utc) if payload.status == "CONFIRMED" else None
    db.commit()
    db.refresh(order)
    return _response(order)
