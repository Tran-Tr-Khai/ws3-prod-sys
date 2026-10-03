from datetime import date
from io import BytesIO
import mimetypes
from pathlib import Path
from uuid import uuid4
from zoneinfo import ZoneInfo

from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile, status
from fastapi.responses import FileResponse, Response, StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import desc, func, select
from sqlalchemy.orm import Session

from app.config.settings import get_settings
from app.db.session import get_db
from app.models.buffing import BuffingCheck, BuffingCheckImage
from app.models.ws3_production import WS3ProductionOrder
from app.schemas.buffing import BuffingCheckCreate, BuffingCheckResponse, BuffingImageResponse
from app.security.auth import get_current_user

router = APIRouter(prefix="/buffing/checks", tags=["buffing"])
MAX_IMAGE_BYTES = 5 * 1024 * 1024
MAX_IMAGES_PER_CHECK = 10
ALLOWED_IMAGE_TYPES = {"image/jpeg", "image/png", "image/webp"}


def media_root() -> Path:
    root = Path(get_settings().media_root).resolve()
    root.mkdir(parents=True, exist_ok=True)
    return root


def image_response(image: BuffingCheckImage) -> BuffingImageResponse:
    return BuffingImageResponse(
        id=image.id,
        original_name=image.original_name,
        mime_type=image.mime_type,
        file_size=image.file_size,
        sort_order=image.sort_order,
        is_primary=image.is_primary,
        url=f"/api/buffing/checks/{image.buffing_check_id}/images/{image.id}/file",
    )


def check_response(check: BuffingCheck) -> BuffingCheckResponse:
    return BuffingCheckResponse(
        id=check.id,
        machine_id=check.machine_id,
        check_date=check.check_date,
        checked_at=check.checked_at,
        operator_name=check.operator_name,
        operator_identifier=check.operator_identifier,
        shift=check.shift,
        order_number=check.order_number,
        check_1=check.check_1,
        check_2=check.check_2,
        check_3=check.check_3,
        check_4=check.check_4,
        check_5=check.check_5,
        remark=check.remark,
        created_at=check.created_at,
        images=[image_response(image) for image in check.images],
    )


@router.post("", response_model=BuffingCheckResponse, status_code=201)
def create_buffing_check(payload: BuffingCheckCreate, db: Session = Depends(get_db)) -> BuffingCheckResponse:
    normalized_order = payload.order_number.strip()
    order_exists = db.scalar(
        select(WS3ProductionOrder.id)
        .where(func.upper(func.trim(WS3ProductionOrder.pk_no)) == normalized_order.upper())
        .limit(1)
    )
    if order_exists is None:
        raise HTTPException(status_code=422, detail="Không tìm thấy mã đơn trong dữ liệu MES / Production order was not found in MES. Check the PKP number and try again.")
    payload.order_number = normalized_order
    check = BuffingCheck(**payload.model_dump())
    db.add(check)
    db.commit()
    db.refresh(check)
    return check_response(check)


@router.get("", response_model=list[BuffingCheckResponse])
def list_buffing_checks(
    check_date: date | None = Query(default=None),
    machine_id: str = Query(default="BU-01", min_length=1, max_length=50),
    db: Session = Depends(get_db),
) -> list[BuffingCheckResponse]:
    statement = select(BuffingCheck).where(BuffingCheck.machine_id == machine_id)
    if check_date is not None:
        statement = statement.where(BuffingCheck.check_date == check_date)
    statement = statement.order_by(desc(BuffingCheck.checked_at), desc(BuffingCheck.id))
    return [check_response(check) for check in db.scalars(statement).all()]


@router.get("/export")
def export_buffing_checks(
    check_date: date = Query(...),
    machine_id: str = Query(default="BU-01", min_length=1, max_length=50),
    language: str = Query(default="vi", pattern="^(vi|en)$"),
    db: Session = Depends(get_db),
    _user=Depends(get_current_user),
) -> StreamingResponse:
    checks = db.scalars(
        select(BuffingCheck)
        .where(BuffingCheck.machine_id == machine_id, BuffingCheck.check_date == check_date)
        .order_by(desc(BuffingCheck.checked_at), desc(BuffingCheck.id))
    ).all()

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Buffing"
    sheet.sheet_view.showGridLines = False
    sheet.merge_cells("A1:L1")
    title = sheet["A1"]
    title.value = "BÁO CÁO CHECKLIST · BUFFING" if language == "vi" else "BUFFING · CHECKLIST REPORT"
    title.font = Font(name="Aptos Display", size=16, bold=True, color="FFFFFF")
    title.fill = PatternFill("solid", fgColor="385466")
    title.alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 32
    sheet.merge_cells("A2:L2")
    metadata = sheet["A2"]
    metadata.value = (
        f"Máy: {machine_id}    |    Ngày ghi: {check_date:%d/%m/%Y}    |    Tổng lượt kiểm tra: {len(checks)}"
        if language == "vi"
        else f"Machine: {machine_id}    |    Record date: {check_date:%d/%m/%Y}    |    Total checks: {len(checks)}"
    )
    metadata.font = Font(name="Aptos", size=10, color="385466")
    metadata.fill = PatternFill("solid", fgColor="EAF0F3")
    metadata.alignment = Alignment(vertical="center")
    sheet.row_dimensions[2].height = 23

    headers = (
        ["THỜI GIAN", "ĐƠN SẢN XUẤT", "NGƯỜI VẬN HÀNH", "ID NHÂN VIÊN", "CA", "KIỂM TRA 1", "KIỂM TRA 2", "KIỂM TRA 3", "KIỂM TRA 4", "KIỂM TRA 5", "GHI CHÚ", "HÌNH ẢNH"]
        if language == "vi"
        else ["TIME", "PRODUCTION ORDER", "OPERATOR", "EMPLOYEE ID", "SHIFT", "CHECK 1", "CHECK 2", "CHECK 3", "CHECK 4", "CHECK 5", "REMARK", "IMAGES"]
    )
    header_row = 4
    for column, value in enumerate(headers, start=1):
        cell = sheet.cell(header_row, column, value)
        cell.font = Font(name="Aptos", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="476271")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=Side(style="medium", color="243B4A"))
    sheet.row_dimensions[header_row].height = 27

    thin_border = Border(bottom=Side(style="thin", color="D5E0E5"))
    for row_index, check in enumerate(checks, start=header_row + 1):
        checked_at = check.checked_at
        if checked_at.tzinfo is not None:
            checked_at = checked_at.astimezone(ZoneInfo("Asia/Ho_Chi_Minh")).replace(tzinfo=None)
        image_names = ", ".join(image.original_name for image in check.images) or "—"
        values = [
            checked_at.strftime("%H:%M"), check.order_number or "—", check.operator_name or "—",
            check.operator_identifier or "—", check.shift or "—", check.check_1, check.check_2,
            check.check_3, check.check_4, check.check_5, check.remark or "—", image_names,
        ]
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row_index, column, value)
            cell.font = Font(name="Aptos", size=10, color="243B4A")
            cell.alignment = Alignment(vertical="center", wrap_text=column in (11, 12))
            cell.border = thin_border
            if row_index % 2:
                cell.fill = PatternFill("solid", fgColor="F3F7F9")
            if 6 <= column <= 10:
                cell.value = ("ĐẠT" if value else "KHÔNG ĐẠT") if language == "vi" else ("PASS" if value else "FAIL")
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.font = Font(name="Aptos", size=9, bold=True, color="247044" if value else "B42318")
                cell.fill = PatternFill("solid", fgColor="E7F4EC" if value else "FDECEC")
        sheet.row_dimensions[row_index].height = 24

    widths = [12, 22, 24, 16, 13, 11, 11, 11, 11, 11, 38, 40]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A5"
    sheet.auto_filter.ref = f"A{header_row}:L{header_row + len(checks)}"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_title_rows = "1:4"

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    filename = f"buffing-report-{check_date.isoformat()}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


def get_check(check_id: int, db: Session) -> BuffingCheck:
    check = db.get(BuffingCheck, check_id)
    if check is None:
        raise HTTPException(status_code=404, detail="Buffing check not found")
    return check


@router.delete("/{check_id}", status_code=status.HTTP_204_NO_CONTENT)
def delete_buffing_check(check_id: int, db: Session = Depends(get_db), _user=Depends(get_current_user)) -> Response:
    check = get_check(check_id, db)
    file_paths = [image.file_path for image in check.images]
    db.delete(check)
    db.commit()

    root = media_root()
    for file_path in file_paths:
        target = (root / file_path).resolve()
        if root in target.parents and target.is_file():
            target.unlink()
    check_directory = (root / str(check_id)).resolve()
    if root in check_directory.parents:
        try:
            check_directory.rmdir()
        except OSError:
            pass
    return Response(status_code=status.HTTP_204_NO_CONTENT)


@router.post("/{check_id}/images", response_model=list[BuffingImageResponse], status_code=status.HTTP_201_CREATED)
async def upload_buffing_images(
    check_id: int,
    files: list[UploadFile] = File(...),
    db: Session = Depends(get_db),
    _user=Depends(get_current_user),
) -> list[BuffingImageResponse]:
    check = get_check(check_id, db)
    if not files or len(check.images) + len(files) > MAX_IMAGES_PER_CHECK:
        raise HTTPException(status_code=400, detail=f"A Buffing check can contain at most {MAX_IMAGES_PER_CHECK} images")

    root = media_root()
    saved: list[BuffingCheckImage] = []
    try:
        for offset, upload in enumerate(files):
            if upload.content_type not in ALLOWED_IMAGE_TYPES:
                raise HTTPException(status_code=415, detail="Only JPG, PNG and WEBP images are supported")
            content = await upload.read()
            if not content or len(content) > MAX_IMAGE_BYTES:
                raise HTTPException(status_code=413, detail="Each image must be smaller than 5 MB")
            extension = mimetypes.guess_extension(upload.content_type) or ".img"
            relative_path = Path(str(check.id)) / f"{uuid4().hex}{extension}"
            target = root / relative_path
            target.parent.mkdir(parents=True, exist_ok=True)
            target.write_bytes(content)
            image = BuffingCheckImage(
                buffing_check_id=check.id,
                file_path=str(relative_path),
                original_name=upload.filename or f"buffing-{check.id}-{offset + 1}{extension}",
                mime_type=upload.content_type,
                file_size=len(content),
                sort_order=len(check.images) + offset,
                is_primary=not check.images and offset == 0,
            )
            db.add(image)
            saved.append(image)
        db.commit()
        for image in saved:
            db.refresh(image)
        return [image_response(image) for image in saved]
    except Exception:
        db.rollback()
        for image in saved:
            target = root / image.file_path
            if target.exists():
                target.unlink()
        raise


@router.get("/{check_id}/images/{image_id}/file", response_class=FileResponse)
def get_buffing_image(check_id: int, image_id: int, db: Session = Depends(get_db), _user=Depends(get_current_user)) -> FileResponse:
    image = db.scalar(select(BuffingCheckImage).where(BuffingCheckImage.id == image_id, BuffingCheckImage.buffing_check_id == check_id))
    if image is None:
        raise HTTPException(status_code=404, detail="Buffing image not found")
    root = media_root()
    target = (root / image.file_path).resolve()
    if root not in target.parents or not target.is_file():
        raise HTTPException(status_code=404, detail="Buffing image file not found")
    return FileResponse(target, media_type=image.mime_type, filename=image.original_name)


@router.patch("/{check_id}/images/{image_id}/primary", response_model=BuffingImageResponse)
def set_primary_buffing_image(check_id: int, image_id: int, db: Session = Depends(get_db), _user=Depends(get_current_user)) -> BuffingImageResponse:
    image = db.scalar(select(BuffingCheckImage).where(BuffingCheckImage.id == image_id, BuffingCheckImage.buffing_check_id == check_id))
    if image is None:
        raise HTTPException(status_code=404, detail="Buffing image not found")
    siblings = db.scalars(select(BuffingCheckImage).where(BuffingCheckImage.buffing_check_id == check_id)).all()
    for sibling in siblings:
        sibling.is_primary = sibling.id == image_id
    db.commit()
    db.refresh(image)
    return image_response(image)
