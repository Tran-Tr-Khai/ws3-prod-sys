"""MES snapshot readers. Header semantics belong here, never in the UI."""
import json
import re
from datetime import datetime
from decimal import Decimal
from io import BytesIO


def text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def key(value: object) -> str:
    return re.sub(r"\s+", "", text(value)).upper()


def read_excel(content: bytes, filename: str) -> list[list[object]]:
    """Expand only actual merged cells; never fill arbitrary blank source cells."""
    if filename.lower().endswith('.xls'):
        import xlrd
        book = xlrd.open_workbook(file_contents=content, formatting_info=True)
        sheet = book.sheet_by_index(0)
        rows = [[xlrd.xldate_as_datetime(c.value, book.datemode) if c.ctype == xlrd.XL_CELL_DATE else c.value for c in sheet.row(i)] for i in range(sheet.nrows)]
        merges = sheet.merged_cells
        for r0, r1, c0, c1 in merges:
            for r in range(r0, r1):
                for c in range(c0, c1):
                    rows[r][c] = rows[r0][c0]
        book.release_resources()
        return rows
    if filename.lower().endswith('.xlsx'):
        from openpyxl import load_workbook
        book = load_workbook(BytesIO(content), data_only=True)
        sheet = book.active
        rows = [list(row) for row in sheet.iter_rows(values_only=True)]
        for merge in sheet.merged_cells.ranges:
            for r in range(merge.min_row - 1, merge.max_row):
                for c in range(merge.min_col - 1, merge.max_col):
                    rows[r][c] = rows[merge.min_row - 1][merge.min_col - 1]
        book.close()
        return rows
    raise ValueError('Use XLS or XLSX files.')


COMMON = {'po_no': 'PO #', 'source_order_no': 'ORDER #', 'item_code': 'ITEM NO', 'item_name': 'FABRIC', 'lot_no': 'LOT #'}
FIELDS = {
    'order': {**COMMON, 'production_date': 'DATE', 'pk_no': 'PK #', 'out_no': 'OUT NO', 'invoice_no': 'INVOICE NO', 'quantity': 'QTY / ROLL'},
    'machine': {**COMMON, 'weaving_date': 'DATE', 'roll_id': 'ROLL #', 'machine_no': 'MACHINE', 'sop_no': 'SOP #', 'out_dyeing_sop': 'OUT DYEING SOP #', 'length_meters': 'MTS'},
    'worker': {**COMMON, 'weaving_date': 'DATE', 'sop_no': 'SOP #', 'machine_no': 'MACHINE', 'shift': 'SHIFT', 'worker': 'WORKER', 'customer': 'CUSTOMER', 'quantity': 'PRODUCTION QTY / ROLL', 'length_meters': 'PRODUCTION QTY / MTS'},
}


def parse_source(rows: list[list[object]], source: str) -> list[dict[str, object]]:
    fields = FIELDS[source]
    required = {key(v) for v in fields.values() if '/' not in v}
    header_index = next((i for i, row in enumerate(rows[:30]) if required.issubset({key(v) for v in row})), None)
    if header_index is None:
        raise ValueError(f'{source}: missing required columns: {", ".join(sorted(required))}')
    headers = [re.sub(r'\s+', ' ', text(v)).upper() for v in rows[header_index]]
    sub = rows[header_index + 1] if header_index + 1 < len(rows) else []
    # MES puts aliases A/B/C in the second row. Only quantity units there
    # replace the group heading; MACHINE-ORDER already has real leaf names.
    if source in ('order', 'worker'):
        group = 'QTY' if source == 'order' else 'PRODUCTION QTY'
        start = next((i for i, h in enumerate(headers) if h == group), None)
        if start is None:
            raise ValueError(f'{source}: missing {group} header')
        for offset, unit in enumerate(('ROLL', 'KGS', 'YDS', 'MTS')):
            index = start + offset
            if index >= len(sub) or key(sub[index]) != unit:
                raise ValueError(f'{source}: invalid {group} unit header: {unit}')
            headers[index] = f'{group} / {unit}'
    positions = {key(h): i for i, h in enumerate(headers) if h}
    result = []
    for line, row in enumerate(rows[header_index + 1:], header_index + 2):
        values = [text(v) for v in row] + [''] * max(0, len(headers) - len(row))
        if not any(values):
            continue
        date_value = values[positions['DATE']]
        if date_value in ('A', 'C', 'DATE'):
            continue
        try:
            date_value = datetime.strptime(date_value[:10], '%Y-%m-%d').date().isoformat() if date_value else None
        except ValueError as exc:
            raise ValueError(f'{source}: invalid DATE at row {line}: {date_value}') from exc
        record = {field: values[positions[key(label)]] or None for field, label in fields.items()}
        record['production_date' if source == 'order' else 'weaving_date'] = date_value
        if 'length_meters' in record:
            record['length_meters'] = Decimal(record['length_meters'].replace(',', '')) if record['length_meters'] else None
            length = record['length_meters']
            if length is not None and (not length.is_finite() or length < 0):
                raise ValueError(f'{source}: invalid length at row {line}')
        if source == 'order':
            quantity = Decimal(record['quantity'] or '0')
            if not quantity.is_finite() or quantity <= 0 or quantity != quantity.to_integral_value():
                raise ValueError(f'{source}: invalid roll count at row {line}')
            if not record['pk_no'] or not record['production_date']:
                raise ValueError(f'{source}: missing PK # or DATE at row {line}')
            mts = Decimal(values[positions[key('QTY / MTS')]].replace(',', '') or '0')
            if not mts.is_finite() or mts < 0:
                raise ValueError(f'{source}: invalid MTS at row {line}')
        if 'machine_no' in record:
            record['machine_no'] = key(record['machine_no']) or None
        identity = 'out_no' if source == 'order' else 'roll_id' if source == 'machine' else 'sop_no'
        if not record[identity]:
            raise ValueError(f'{source}: missing {identity} at row {line}')
        record['raw_data_json'] = json.dumps({'columns': headers, 'row': values}, ensure_ascii=False)
        result.append(record)
    return result
